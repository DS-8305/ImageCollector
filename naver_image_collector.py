from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import threading
import time
from collections.abc import Callable, Iterable
from dataclasses import asdict, dataclass
from io import BytesIO
from pathlib import Path
from urllib.parse import quote_plus

import requests

try:
    from PIL import Image
except ImportError:
    Image = None

try:
    from openpyxl import load_workbook
except ImportError:
    load_workbook = None


METADATA_FIELDS = [
    "keyword",
    "index",
    "file_name",
    "thumbnail_url",
    "source_url",
    "title",
    "source",
    "status",
    "error",
]
IMAGE_FILE_PATTERN = re.compile(r"^image_\d+\.(jpg|jpeg|png|webp)$", re.IGNORECASE)
USER_AGENT = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
    "AppleWebKit/537.36 (KHTML, like Gecko) "
    "Chrome/136.0.0.0 Safari/537.36"
)
DEFAULT_LIMIT = 50
DEFAULT_OUTPUT_DIR = "downloaded_images"
DEFAULT_MIN_THUMBNAIL_SIZE = 150


@dataclass(slots=True)
class ImageRecord:
    keyword: str
    index: int
    file_name: str
    thumbnail_url: str
    source_url: str
    title: str
    source: str
    status: str
    error: str


@dataclass(slots=True)
class ImageCandidate:
    thumbnail_url: str
    source_url: str
    title: str
    source: str


@dataclass(frozen=True, slots=True)
class BatchResult:
    keyword: str
    saved_count: int
    status: str
    error: str


def safe_folder_name(value: str) -> str:
    safe_value = re.sub(r"[^\w\-]+", "_", value.strip(), flags=re.UNICODE)
    safe_value = safe_value.strip("_")
    if safe_value:
        return safe_value[:120]
    return hashlib.sha1(value.encode("utf-8")).hexdigest()[:10]


def ensure_keyword_folder(output_root: str | Path, keyword: str) -> Path:
    folder_path = Path(output_root) / safe_folder_name(keyword)
    folder_path.mkdir(parents=True, exist_ok=True)
    return folder_path


def write_metadata(folder_path: str | Path, records: list[ImageRecord]) -> None:
    folder = Path(folder_path)
    folder.mkdir(parents=True, exist_ok=True)
    rows = [asdict(record) for record in records]
    with (folder / "metadata.csv").open("w", encoding="utf-8-sig", newline="") as file:
        writer = csv.DictWriter(file, fieldnames=METADATA_FIELDS, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)
    with (folder / "metadata.json").open("w", encoding="utf-8") as file:
        json.dump(rows, file, ensure_ascii=False, indent=2)


def clear_generated_images(folder_path: str | Path) -> None:
    folder = Path(folder_path)
    if not folder.exists():
        return
    for child in folder.iterdir():
        if child.is_file() and IMAGE_FILE_PATTERN.match(child.name):
            child.unlink()


def load_keywords(file_path: str | Path) -> list[str]:
    path = Path(file_path)
    suffix = path.suffix.lower()
    if suffix == ".csv":
        return _load_keywords_from_csv(path)
    if suffix == ".xlsx":
        return _load_keywords_from_xlsx(path)
    raise ValueError("Only .csv and .xlsx files are supported.")


def _dedupe_keywords(values: list[str]) -> list[str]:
    keywords: list[str] = []
    seen: set[str] = set()
    for value in values:
        keyword = value.strip()
        if not keyword or keyword in seen:
            continue
        seen.add(keyword)
        keywords.append(keyword)
    if not keywords:
        raise ValueError("No non-empty keyword rows were found.")
    return keywords


def _load_keywords_from_csv(path: Path) -> list[str]:
    with path.open("r", encoding="utf-8-sig", newline="") as file:
        reader = csv.DictReader(file)
        headers = {header.strip().lower(): header for header in (reader.fieldnames or []) if header}
        if "keyword" not in headers:
            raise ValueError("CSV header must include keyword.")
        keyword_header = headers["keyword"]
        return _dedupe_keywords([(row.get(keyword_header) or "") for row in reader])


def _load_keywords_from_xlsx(path: Path) -> list[str]:
    if load_workbook is None:
        raise RuntimeError("openpyxl is required for XLSX files.")
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        rows = list(worksheet.iter_rows(values_only=True))
    finally:
        workbook.close()
    if not rows:
        raise ValueError("XLSX file is empty.")
    headers = [str(cell).strip().lower() if cell is not None else "" for cell in rows[0]]
    try:
        keyword_index = headers.index("keyword")
    except ValueError as exc:
        raise ValueError("XLSX header must include keyword.") from exc
    return _dedupe_keywords(
        [
            str(row[keyword_index]) if keyword_index < len(row) and row[keyword_index] is not None else ""
            for row in rows[1:]
        ]
    )


class NaverImageCollector:
    def __init__(
        self,
        *,
        output_root: str | Path = DEFAULT_OUTPUT_DIR,
        headless: bool = True,
        timeout_seconds: int = 20,
        min_thumbnail_size: int = DEFAULT_MIN_THUMBNAIL_SIZE,
        logger: Callable[[str], None] | None = None,
        driver=None,
    ) -> None:
        self.output_root = Path(output_root)
        self.headless = headless
        self.timeout_seconds = timeout_seconds
        self.min_thumbnail_size = min_thumbnail_size
        self.logger = logger
        self.session = requests.Session()
        self.session.headers.update({"User-Agent": USER_AGENT, "Referer": "https://search.naver.com/"})
        self._owns_driver = driver is None
        self.driver = driver or self._build_driver()

    def _build_driver(self):
        try:
            from selenium import webdriver
            from selenium.webdriver.chrome.options import Options
            from selenium.webdriver.chrome.service import Service
            from webdriver_manager.chrome import ChromeDriverManager
        except ImportError as exc:
            raise RuntimeError(
                "Selenium dependencies are missing. Run: python -m pip install -r requirements.txt"
            ) from exc

        chrome_options = Options()
        if self.headless:
            chrome_options.add_argument("--headless=new")
        chrome_options.add_argument("--no-sandbox")
        chrome_options.add_argument("--disable-dev-shm-usage")
        chrome_options.add_argument("--window-size=1400,2400")
        chrome_options.add_argument(f"user-agent={USER_AGENT}")
        service = Service(ChromeDriverManager().install())
        return webdriver.Chrome(service=service, options=chrome_options)

    def close(self) -> None:
        if not self._owns_driver:
            return
        try:
            self.driver.quit()
        except Exception:
            pass

    def collect_keyword(self, keyword: str, *, limit: int = DEFAULT_LIMIT) -> list[Path]:
        keyword = keyword.strip()
        if not keyword:
            raise ValueError("Keyword is empty.")
        if limit < 1:
            raise ValueError("Limit must be at least 1.")

        output_folder = ensure_keyword_folder(self.output_root, keyword)
        clear_generated_images(output_folder)
        self._log(f"Search started: {keyword}")
        self.driver.get(self._search_url(keyword))
        time.sleep(1.5)

        candidates = self._collect_candidates(limit)
        if not candidates:
            raise RuntimeError("No image candidates were found. Naver page selectors may need updating.")

        records: list[ImageRecord] = []
        saved_paths: list[Path] = []
        seen_urls: set[str] = set()
        saved_index = 1

        for candidate in candidates:
            if len(saved_paths) >= limit:
                break
            if candidate.thumbnail_url in seen_urls:
                continue
            seen_urls.add(candidate.thumbnail_url)
            file_name = f"image_{saved_index:02d}.png"
            file_path = output_folder / file_name
            try:
                self._download_thumbnail(candidate.thumbnail_url, file_path)
                status = "saved"
                error = ""
                saved_paths.append(file_path)
                self._log(f"Saved: {file_name}")
                saved_index += 1
            except Exception as exc:
                status = "failed"
                error = str(exc)
                self._log(f"Save failed: {candidate.thumbnail_url} / {exc}")

            records.append(
                ImageRecord(
                    keyword=keyword,
                    index=len(records) + 1,
                    file_name=file_name if status == "saved" else "",
                    thumbnail_url=candidate.thumbnail_url,
                    source_url=candidate.source_url,
                    title=candidate.title,
                    source=candidate.source,
                    status=status,
                    error=error,
                )
            )

        if len(saved_paths) < limit:
            self._log(f"Warning: requested {limit} images, saved {len(saved_paths)}.")
        write_metadata(output_folder, records)
        self._log(f"Search complete: {keyword} / {len(saved_paths)} images")
        return saved_paths

    def _search_url(self, keyword: str) -> str:
        return f"https://search.naver.com/search.naver?where=image&sm=tab_opt&sort=date&query={quote_plus(keyword)}"

    def _collect_candidates(self, limit: int) -> list[ImageCandidate]:
        candidates: list[ImageCandidate] = []
        seen: set[str] = set()
        stale_rounds = 0
        target = max(limit * 2, limit + 10)

        for _ in range(14):
            for candidate in self._extract_candidates_from_page():
                if candidate.thumbnail_url in seen:
                    continue
                seen.add(candidate.thumbnail_url)
                candidates.append(candidate)
                if len(candidates) >= target:
                    return candidates

            before = len(candidates)
            self.driver.execute_script("window.scrollTo(0, document.body.scrollHeight);")
            time.sleep(1.0)
            after = len(candidates)
            stale_rounds = stale_rounds + 1 if after == before else 0
            if stale_rounds >= 3 and len(candidates) >= limit:
                break
        return candidates

    def _extract_candidates_from_page(self) -> list[ImageCandidate]:
        raw_items = self.driver.execute_script(
            """
            const results = [];
            const images = Array.from(document.querySelectorAll('img'));
            for (const img of images) {
              const src = img.currentSrc || img.src || img.getAttribute('data-src') || '';
              if (!src || src.startsWith('data:')) continue;
              const box = img.getBoundingClientRect();
              if (box.width < 80 || box.height < 80) continue;
              const anchor = img.closest('a');
              const container = img.closest('div, li, section') || img.parentElement;
              const text = container ? container.innerText : '';
              results.push({
                thumbnailUrl: src,
                sourceUrl: anchor ? anchor.href : '',
                title: img.alt || img.title || '',
                source: text.split('\\n').slice(0, 2).join(' / '),
                width: Math.round(box.width),
                height: Math.round(box.height)
              });
            }
            return results;
            """
        )
        candidates: list[ImageCandidate] = []
        for item in raw_items or []:
            thumbnail_url = str(item.get("thumbnailUrl") or "").strip()
            if not thumbnail_url or thumbnail_url.startswith("data:"):
                continue
            candidates.append(
                ImageCandidate(
                    thumbnail_url=thumbnail_url,
                    source_url=str(item.get("sourceUrl") or "").strip(),
                    title=str(item.get("title") or "").strip(),
                    source=str(item.get("source") or "").strip(),
                )
            )
        return candidates

    def _download_thumbnail(self, url: str, file_path: Path) -> None:
        response = self.session.get(url, timeout=20)
        response.raise_for_status()
        content = response.content
        if Image is not None:
            with Image.open(BytesIO(content)) as image:
                if image.width < self.min_thumbnail_size or image.height < self.min_thumbnail_size:
                    raise ValueError(
                        f"Image {image.width}x{image.height} is smaller than minimum "
                        f"{self.min_thumbnail_size}x{self.min_thumbnail_size}."
                    )
                image.convert("RGBA").save(file_path, "PNG", optimize=True)
            return
        file_path.write_bytes(content)

    def _log(self, message: str) -> None:
        if self.logger:
            self.logger(message)
        else:
            print(message)


def collect_single_keyword(
    keyword: str,
    *,
    limit: int = DEFAULT_LIMIT,
    output_root: str | Path = DEFAULT_OUTPUT_DIR,
    headless: bool = True,
    min_thumbnail_size: int = DEFAULT_MIN_THUMBNAIL_SIZE,
    logger: Callable[[str], None] | None = None,
) -> list[Path]:
    collector = NaverImageCollector(
        output_root=output_root,
        headless=headless,
        min_thumbnail_size=min_thumbnail_size,
        logger=logger,
    )
    try:
        return collector.collect_keyword(keyword, limit=limit)
    finally:
        collector.close()


def run_batch_collection(
    keywords: Iterable[str],
    *,
    limit: int = DEFAULT_LIMIT,
    output_root: str | Path = DEFAULT_OUTPUT_DIR,
    headless: bool = True,
    min_thumbnail_size: int = DEFAULT_MIN_THUMBNAIL_SIZE,
    logger: Callable[[str], None] | None = None,
    collector_func: Callable[[str, int, Path, bool, int], list[Path]] | None = None,
) -> list[BatchResult]:
    results: list[BatchResult] = []
    output_path = Path(output_root)
    collect = collector_func or _collect_single_keyword_for_batch
    keyword_list = [keyword.strip() for keyword in keywords if keyword.strip()]
    for index, keyword in enumerate(keyword_list, start=1):
        _log(logger, f"[{index}/{len(keyword_list)}] Start: {keyword}")
        try:
            saved_paths = collect(keyword, limit, output_path, headless, min_thumbnail_size)
            results.append(BatchResult(keyword=keyword, saved_count=len(saved_paths), status="success", error=""))
            _log(logger, f"[{index}/{len(keyword_list)}] Done: {keyword} / {len(saved_paths)} images")
        except Exception as exc:
            results.append(BatchResult(keyword=keyword, saved_count=0, status="failed", error=str(exc)))
            _log(logger, f"[{index}/{len(keyword_list)}] Failed: {keyword} / {exc}")
    return results


def _collect_single_keyword_for_batch(
    keyword: str,
    limit: int,
    output_root: Path,
    headless: bool,
    min_thumbnail_size: int,
) -> list[Path]:
    return collect_single_keyword(
        keyword,
        limit=limit,
        output_root=output_root,
        headless=headless,
        min_thumbnail_size=min_thumbnail_size,
    )


def _log(logger: Callable[[str], None] | None, message: str) -> None:
    if logger:
        logger(message)
    else:
        print(message)


class NaverImageCollectorGUI:
    def __init__(self, root) -> None:
        import tkinter as tk
        from tkinter import ttk

        self.tk = tk
        self.ttk = ttk
        self.root = root
        self.root.title("Naver Image Collector")
        self.root.geometry("900x700")
        self.keyword_var = tk.StringVar(value="")
        self.limit_var = tk.StringVar(value=str(DEFAULT_LIMIT))
        self.min_thumbnail_size_var = tk.StringVar(value=str(DEFAULT_MIN_THUMBNAIL_SIZE))
        self.batch_file_var = tk.StringVar(value="")
        self.batch_keywords: list[str] = []
        self.validated_batch_file = ""
        self.output_dir_var = tk.StringVar(value=DEFAULT_OUTPUT_DIR)
        self.show_browser_var = tk.BooleanVar(value=False)
        self.run_buttons = []
        self._build_ui()

    def _build_ui(self) -> None:
        tk = self.tk
        ttk = self.ttk
        frame = ttk.Frame(self.root, padding=16)
        frame.pack(fill=tk.BOTH, expand=True)
        frame.columnconfigure(0, weight=1)

        single_frame = ttk.LabelFrame(frame, text="Single keyword", padding=12)
        single_frame.grid(row=0, column=0, sticky="ew")
        single_frame.columnconfigure(1, weight=1)

        ttk.Label(single_frame, text="Keyword").grid(row=0, column=0, sticky=tk.W, padx=(0, 8), pady=4)
        ttk.Entry(single_frame, textvariable=self.keyword_var).grid(row=0, column=1, sticky="ew", pady=4)
        single_button = ttk.Button(single_frame, text="Run search", command=self._start_single_run)
        single_button.grid(row=0, column=2, padx=(8, 0), pady=4)
        self.run_buttons.append(single_button)

        options_frame = ttk.Frame(frame)
        options_frame.grid(row=1, column=0, sticky="ew", pady=(12, 0))
        options_frame.columnconfigure(1, weight=1)
        ttk.Label(options_frame, text="Limit").grid(row=0, column=0, sticky=tk.W, padx=(0, 8), pady=4)
        ttk.Entry(options_frame, textvariable=self.limit_var, width=8).grid(row=0, column=1, sticky=tk.W, pady=4)
        ttk.Label(options_frame, text="Min size").grid(row=0, column=2, sticky=tk.W, padx=(12, 8), pady=4)
        ttk.Entry(options_frame, textvariable=self.min_thumbnail_size_var, width=8).grid(
            row=0, column=3, sticky=tk.W, pady=4
        )
        ttk.Checkbutton(options_frame, text="Show browser", variable=self.show_browser_var).grid(
            row=0, column=4, sticky=tk.W, padx=(12, 0), pady=4
        )

        output_frame = ttk.LabelFrame(frame, text="Output folder", padding=12)
        output_frame.grid(row=2, column=0, sticky="ew", pady=(12, 0))
        output_frame.columnconfigure(0, weight=1)
        ttk.Entry(output_frame, textvariable=self.output_dir_var).grid(row=0, column=0, sticky="ew")
        ttk.Button(output_frame, text="Choose", command=self._choose_output_dir).grid(row=0, column=1, padx=(8, 0))

        batch_frame = ttk.LabelFrame(frame, text="CSV / XLSX batch", padding=12)
        batch_frame.grid(row=3, column=0, sticky="ew", pady=(12, 0))
        batch_frame.columnconfigure(0, weight=1)
        ttk.Entry(batch_frame, textvariable=self.batch_file_var).grid(row=0, column=0, sticky="ew")
        ttk.Button(batch_frame, text="Upload CSV", command=self._choose_batch_file).grid(row=0, column=1, padx=(8, 0))
        batch_button = ttk.Button(batch_frame, text="Search images", command=self._start_batch_run)
        batch_button.grid(row=0, column=2, padx=(8, 0))
        self.run_buttons.append(batch_button)

        ttk.Label(frame, text="Log").grid(row=4, column=0, sticky=tk.W, pady=(12, 4))
        self.log_text = tk.Text(frame, wrap=tk.WORD, font=("Consolas", 10), height=18)
        self.log_text.grid(row=5, column=0, sticky="nsew")
        frame.rowconfigure(5, weight=1)
        self._append_log("Ready. Enter a keyword or choose a CSV/XLSX file.")

    def _choose_batch_file(self) -> None:
        from tkinter import filedialog, messagebox

        file_path = filedialog.askopenfilename(
            title="Choose CSV or XLSX",
            filetypes=[("Batch files", "*.csv;*.xlsx"), ("All files", "*.*")],
        )
        if file_path:
            try:
                self._validate_batch_file(file_path)
            except (OSError, ValueError, RuntimeError) as exc:
                messagebox.showerror("CSV error", str(exc))

    def _validate_batch_file(self, file_path: str) -> list[str]:
        keywords = load_keywords(file_path)
        self.batch_file_var.set(file_path)
        self.batch_keywords = keywords
        self.validated_batch_file = file_path
        self._append_log(f"CSV loaded: {len(keywords)} keywords / first: {keywords[0]}")
        return keywords

    def _choose_output_dir(self) -> None:
        from tkinter import filedialog

        folder_path = filedialog.askdirectory(title="Choose output folder")
        if folder_path:
            self.output_dir_var.set(folder_path)

    def _parse_limit(self) -> int:
        try:
            limit = int(self.limit_var.get().strip())
        except ValueError as exc:
            raise ValueError("Limit must be a number.") from exc
        if limit < 1:
            raise ValueError("Limit must be at least 1.")
        return limit

    def _parse_min_thumbnail_size(self) -> int:
        try:
            min_size = int(self.min_thumbnail_size_var.get().strip())
        except ValueError as exc:
            raise ValueError("Min size must be a number.") from exc
        if min_size < 1:
            raise ValueError("Min size must be at least 1.")
        return min_size

    def _start_single_run(self) -> None:
        from tkinter import messagebox

        keyword = self.keyword_var.get().strip()
        if not keyword:
            messagebox.showerror("Input error", "Enter a keyword.")
            return
        try:
            limit = self._parse_limit()
            min_thumbnail_size = self._parse_min_thumbnail_size()
        except ValueError as exc:
            messagebox.showerror("Input error", str(exc))
            return
        self._run_in_thread(self._run_single_task, keyword, limit, min_thumbnail_size)

    def _start_batch_run(self) -> None:
        from tkinter import messagebox

        batch_file = self.batch_file_var.get().strip()
        if not batch_file:
            messagebox.showerror("Input error", "Choose a CSV or XLSX file.")
            return
        try:
            limit = self._parse_limit()
            min_thumbnail_size = self._parse_min_thumbnail_size()
            if batch_file != self.validated_batch_file:
                self._validate_batch_file(batch_file)
        except (OSError, ValueError, RuntimeError) as exc:
            messagebox.showerror("Input error", str(exc))
            return
        self._run_in_thread(self._run_batch_task, list(self.batch_keywords), limit, min_thumbnail_size)

    def _run_in_thread(self, target, *args) -> None:
        self._set_running(True)
        thread = threading.Thread(target=self._thread_wrapper, args=(target, args), daemon=True)
        thread.start()

    def _thread_wrapper(self, target, args) -> None:
        try:
            target(*args)
        finally:
            self.root.after(0, self._set_running, False)

    def _run_single_task(self, keyword: str, limit: int, min_thumbnail_size: int) -> None:
        try:
            saved_paths = collect_single_keyword(
                keyword,
                limit=limit,
                output_root=self.output_dir_var.get().strip() or DEFAULT_OUTPUT_DIR,
                headless=not self.show_browser_var.get(),
                min_thumbnail_size=min_thumbnail_size,
                logger=self._append_log,
            )
            self._append_log(f"Done: {keyword} / {len(saved_paths)} images")
        except Exception as exc:
            self._append_log(f"Failed: {keyword} / {exc}")

    def _run_batch_task(self, keywords: list[str], limit: int, min_thumbnail_size: int) -> None:
        try:
            results = run_batch_collection(
                keywords,
                limit=limit,
                output_root=self.output_dir_var.get().strip() or DEFAULT_OUTPUT_DIR,
                headless=not self.show_browser_var.get(),
                min_thumbnail_size=min_thumbnail_size,
                logger=self._append_log,
            )
            succeeded = sum(1 for result in results if result.status == "success")
            failed = sum(1 for result in results if result.status == "failed")
            self._append_log(f"Batch done: success {succeeded}, failed {failed}")
        except Exception as exc:
            self._append_log(f"Batch failed: {exc}")

    def _set_running(self, running: bool) -> None:
        state = self.tk.DISABLED if running else self.tk.NORMAL
        for button in self.run_buttons:
            button.configure(state=state)

    def _append_log(self, message: str) -> None:
        self.root.after(0, self._write_log, message)

    def _write_log(self, message: str) -> None:
        self.log_text.insert(self.tk.END, f"{message}\n")
        self.log_text.see(self.tk.END)


def launch_gui() -> None:
    import tkinter as tk

    root = tk.Tk()
    NaverImageCollectorGUI(root)
    root.mainloop()


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Search Naver Images and download thumbnails.")
    parser.add_argument("--keyword", help="Keyword to search.")
    parser.add_argument("--batch-file", help="CSV or XLSX file with a keyword column.")
    parser.add_argument("--limit", type=int, default=DEFAULT_LIMIT, help="Images to save per keyword.")
    parser.add_argument(
        "--min-thumbnail-size",
        type=int,
        default=DEFAULT_MIN_THUMBNAIL_SIZE,
        help="Minimum width and height in pixels for saved images.",
    )
    parser.add_argument("--output-dir", default=DEFAULT_OUTPUT_DIR, help="Output directory.")
    parser.add_argument("--show-browser", action="store_true", help="Show the Chrome browser window.")
    parser.add_argument("--gui", action="store_true", help="Open the desktop GUI.")
    return parser


def main() -> None:
    parser = build_parser()
    args = parser.parse_args()

    if args.gui or (not args.keyword and not args.batch_file):
        launch_gui()
        return

    if args.limit < 1:
        parser.error("--limit must be at least 1.")
    if args.min_thumbnail_size < 1:
        parser.error("--min-thumbnail-size must be at least 1.")

    if args.batch_file:
        keywords = load_keywords(args.batch_file)
        results = run_batch_collection(
            keywords,
            limit=args.limit,
            output_root=args.output_dir,
            headless=not args.show_browser,
            min_thumbnail_size=args.min_thumbnail_size,
        )
        succeeded = sum(1 for result in results if result.status == "success")
        failed = sum(1 for result in results if result.status == "failed")
        print(f"Batch complete: success {succeeded}, failed {failed}")
        return

    saved_paths = collect_single_keyword(
        args.keyword,
        limit=args.limit,
        output_root=args.output_dir,
        headless=not args.show_browser,
        min_thumbnail_size=args.min_thumbnail_size,
    )
    print(f"Saved {len(saved_paths)} images.")


if __name__ == "__main__":
    main()
